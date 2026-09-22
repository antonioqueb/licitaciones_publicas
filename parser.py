"""Pure, side-effect-free ComprasMX reader. Also exercised without an Odoo server."""
import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from fnmatch import fnmatchcase
from datetime import date, datetime, timezone
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import openpyxl


class ImportValidationError(ValueError):
    pass


IDENTIFIER = re.compile(r'(?<![A-Z0-9])([A-Z]{2}-\d+-[A-Z0-9]+-[A-Z0-9]+-[A-Z]-\d+-\d{4})(?![A-Z0-9])', re.I)
DATE_FIELDS = ('fecha_junta_aclaraciones', 'fecha_limite_preguntas',
               'fecha_entrega_muestras', 'fecha_apertura', 'fecha_fallo')
PARTIDA_FIELDS = ('numero', 'partida_especifica', 'clave_cucop', 'descripcion_cucop',
                  'descripcion_detallada', 'unidad_medida', 'cantidad', 'cantidad_min', 'cantidad_max', 'cantidad_pendiente')
LIST_METADATA_FIELDS = ('numero_listado', 'caracter_publicado', 'siglas_dependencia', 'tipo_publicacion')
PORTAL_FIELDS = ('identificador', 'nombre_publicado', 'codigo_expediente',
                 'unidad_codigo', 'unidad_nombre', 'entidad_codigo', 'entidad_nombre', 'estatus', 'tipo_codigo') + DATE_FIELDS + LIST_METADATA_FIELDS


def norm(value):
    text = unicodedata.normalize('NFKD', str(value if value is not None else ''))
    return re.sub(r'[^a-z0-9]+', ' ', ''.join(c for c in text if not unicodedata.combining(c)).lower()).strip()


# Los nombres y alias de columnas pertenecen al catálogo, no al lector.
MAPPABLE_FIELDS = set(PORTAL_FIELDS) | set(PARTIDA_FIELDS) | {'codigo_sai', 'nombre_sai'}
DETAIL_TYPES = {'detalle_bienes', 'detalle_servicios', 'detalle_rangos'}


def header_norm(value):
    """Comparación Unicode sin acentos, signos, diferencias de caja o espacios."""
    return norm(value).upper()


def filename_norm(value):
    text = unicodedata.normalize('NFKD', str(value or ''))
    return re.sub(r'\s+', ' ', ''.join(c for c in text if not unicodedata.combining(c))).upper().strip()


def profile_config(values):
    """Validate user configuration once; never eval Python supplied in a field."""
    result = dict(values)
    for key in ('hojas', 'fila_encabezado', 'total_columnas'):
        value = result.get(key)
        if isinstance(value, bool) or not isinstance(value, int) or value < 1:
            raise ImportValidationError('%s debe ser un entero positivo.' % key)
    def decode(key, default):
        raw = result.get(key)
        try:
            return json.loads(raw) if isinstance(raw, str) and raw.strip() else (raw or default)
        except (ValueError, TypeError) as exc:
            raise ImportValidationError('%s debe contener JSON válido.' % key) from exc
    words = decode('palabras_clave_deteccion', [])
    expected = decode('encabezados_esperados', {'cols': []})
    mapping = decode('mapeo_columnas', {})
    if not isinstance(words, list) or any(not isinstance(w, str) or not header_norm(w) for w in words):
        raise ImportValidationError('Palabras clave debe ser un array JSON de textos no vacíos.')
    if not isinstance(expected, dict) or set(expected) - {'cols', 'columnas_alternativas'}:
        raise ImportValidationError('Encabezados esperados debe contener cols y, opcionalmente, columnas_alternativas.')
    cols = expected.get('cols', [])
    if not isinstance(cols, list) or any(not isinstance(c, str) or not header_norm(c) for c in cols):
        raise ImportValidationError('cols debe ser un array de encabezados no vacíos.')
    if len({header_norm(c) for c in cols}) != len(cols):
        raise ImportValidationError('Hay encabezados esperados duplicados o equivalentes.')
    alternatives = expected.get('columnas_alternativas', [])
    if not isinstance(alternatives, list) or any(type(n) is not int or n < 1 for n in alternatives):
        raise ImportValidationError('columnas_alternativas debe contener enteros positivos.')
    if not isinstance(mapping, dict) or set(mapping) - MAPPABLE_FIELDS:
        raise ImportValidationError('El mapeo contiene campos destino no permitidos.')
    header_map = {}
    for field, aliases in mapping.items():
        if not isinstance(aliases, list) or any(not isinstance(a, str) or not header_norm(a) for a in aliases):
            raise ImportValidationError('Cada campo del mapeo debe contener un array de nombres de columna.')
        for alias in aliases:
            key = header_norm(alias)
            if key in header_map and header_map[key] != field:
                raise ImportValidationError('Un mismo encabezado no puede mapearse a dos campos: %s.' % alias)
            header_map[key] = field
    if not words and not (result.get('patron_nombre') or '').strip():
        raise ImportValidationError('Indique palabras clave o al menos un patrón de nombre.')
    result.update(palabras_clave_deteccion=words, encabezados_esperados=expected,
                  mapeo_columnas=mapping, header_map=header_map)
    return result


DETAIL_REQUIRED = {'numero', 'partida_especifica', 'clave_cucop', 'descripcion_detallada', 'unidad_medida'}
LIST_REQUIRED = {'identificador', 'nombre_publicado', 'unidad_nombre', 'estatus', 'tipo_codigo'}
TYPE_CODES = {'adq': 'ADQ', 'adquisiciones': 'ADQ', 'adquisicion': 'ADQ',
              'servicios': 'SER', 'servicio': 'SER', 'ser': 'SER',
              'obra publica': 'OBR', 'obras publicas': 'OBR', 'obr': 'OBR',
              'arrendamiento': 'ARR', 'arrendamientos': 'ARR', 'arr': 'ARR'}


def identifier_parts(identifier):
    match = IDENTIFIER.fullmatch(str(identifier or '').strip())
    if not match:
        raise ImportValidationError('Identificador de procedimiento inválido: %s' % identifier)
    parts = match.group(1).upper().split('-')
    return dict(tipo_procedimiento=parts[0], caracter=parts[-3], consecutivo=parts[-2],
                ejercicio=parts[-1])


def business_key(row):
    # JSON preserves boundaries: no collisions from a literal separator in a description.
    triple = [row.get(k) or '' for k in ('partida_especifica', 'clave_cucop', 'descripcion_detallada')]
    return hashlib.md5(json.dumps(triple, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()


def digest(payload):
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()


def diff_rows(incoming, current, previous_keys=(), key='identificador'):
    """Compare data without touching internal decisions. Gone = previous snapshot only."""
    result = {k: [] for k in ('nuevos', 'cambios', 'sin_cambios', 'gone')}
    seen = set()
    for row in incoming:
        identity = row[key]
        if identity in seen:
            raise ImportValidationError('Clave de negocio duplicada: %s' % identity)
        seen.add(identity)
        old = current.get(identity)
        if old is None:
            result['nuevos'].append({'key': identity, 'values': row})
            continue
        delta = {k: {'antes': old.get(k), 'despues': v} for k, v in row.items() if old.get(k) != v}
        entry = {'key': identity, 'values': row, 'diff': delta}
        result['cambios' if delta else 'sin_cambios'].append(entry)
    result['gone'] = [{'key': k, 'values': current[k]} for k in sorted(set(previous_keys) - seen) if k in current]
    return result


def code(value):
    if isinstance(value, bool):
        raise ImportValidationError('Un código no puede ser booleano.')
    if isinstance(value, (float, int)):
        if not math.isfinite(value) or int(value) != value:
            raise ImportValidationError('Código numérico no entero: %s' % value)
        return str(int(value))
    return str(value if value is not None else '').strip()


def number(value, required=False):
    if value is None or value == '':
        if required:
            raise ImportValidationError('Falta una cantidad obligatoria o una fórmula no tiene valor calculado.')
        return 0.0
    try:
        if isinstance(value, bool):
            raise ValueError()
        parsed = float(str(value).replace(',', '').strip())
        if not math.isfinite(parsed) or parsed < 0:
            raise ValueError()
        return parsed
    except (TypeError, ValueError) as exc:
        raise ImportValidationError('Cantidad inválida: %s' % value) from exc


def portal_datetime(value, tz_name):
    if value is None or value == '':
        return False
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        parsed = datetime.combine(value, datetime.min.time())
    else:
        parsed = None
        for fmt in ('%d/%m/%Y %H:%M:%S', '%d/%m/%Y %H:%M', '%d/%m/%Y',
                    '%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
            try:
                parsed = datetime.strptime(str(value).strip(), fmt)
                break
            except ValueError:
                continue
        if parsed is None:
            raise ImportValidationError('Fecha inválida: %s. Use dd/mm/yyyy HH:MM.' % value)
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=ZoneInfo(tz_name))
    return parsed.astimezone(timezone.utc).replace(tzinfo=None).strftime('%Y-%m-%d %H:%M:%S')


class WorkbookReader:
    MAX_BYTES = 25 * 1024 * 1024
    MAX_UNCOMPRESSED = 150 * 1024 * 1024
    MAX_ROWS = 50000

    def __init__(self, binary_data, filename, tz_name='America/Mexico_City', *, profiles=None):
        self.filename = filename
        self.tz_name = tz_name
        try:
            ZoneInfo(tz_name)
        except (ZoneInfoNotFoundError, ValueError) as exc:
            raise ImportValidationError('Zona horaria inválida.') from exc
        if not filename.lower().endswith('.xlsx') or len(binary_data) > self.MAX_BYTES:
            raise ImportValidationError('Seleccione un archivo .xlsx de hasta 25 MB.')
        try:
            with zipfile.ZipFile(io.BytesIO(binary_data)) as archive:
                if sum(z.file_size for z in archive.infolist()) > self.MAX_UNCOMPRESSED:
                    raise ImportValidationError('El Excel supera el tamaño descomprimido permitido.')
            self.wb = openpyxl.load_workbook(io.BytesIO(binary_data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, OSError, KeyError, ValueError) as exc:
            raise ImportValidationError('El archivo no es un Excel válido: %s' % exc) from exc
        self.identificadores_contenido = set()
        try:
            self._detect(profiles or [])
        except Exception:
            self.close()
            raise

    def _detect(self, profiles):
        errors = []
        candidates = []
        for raw in profiles:
            if not raw.get('activo', True):
                continue
            config = profile_config(raw)
            sheet_name = config.get('hoja_nombre')
            sheets = [ws for ws in self.wb.worksheets if not sheet_name or header_norm(ws.title) == header_norm(sheet_name)]
            patterns = [filename_norm(p) for p in (config.get('patron_nombre') or '').split(',') if p.strip()]
            filename_match = any(fnmatchcase(filename_norm(self.filename), p) for p in patterns)
            if not sheets:
                if filename_match:
                    errors.append('%s: no existe la hoja %s.' % (config['name'], sheet_name))
                continue
            ws = sheets[0]
            row = next(ws.iter_rows(min_row=config['fila_encabezado'], max_row=config['fila_encabezado'], values_only=True), ())
            values = list(row)
            while values and values[-1] in (None, ''):
                values.pop()
            names = {header_norm(v) for v in values if v is not None}
            words = {header_norm(w) for w in config['palabras_clave_deteccion']}
            keyword_match = bool(words) and words <= names
            if not keyword_match and not filename_match:
                continue
            headers = [config['header_map'].get(header_norm(v)) for v in values]
            mapped = set(headers) - {None}
            kind = config['tipo_dato']
            if kind in DETAIL_TYPES:
                subtype = ('detalle_rangos' if {'cantidad_min', 'cantidad_max'} <= mapped else
                           'detalle_bienes' if 'cantidad' in mapped else 'detalle_servicios')
                if subtype != kind:
                    continue
            expected = {header_norm(v) for v in config['encabezados_esperados'].get('cols', [])}
            match = 100 if not expected else 100 * len(expected & names) / len(expected)
            problem = None
            if len(self.wb.worksheets) != config['hojas']:
                problem = 'se esperaban %s hojas; se encontraron %s' % (config['hojas'], len(self.wb.worksheets))
            elif len(values) not in [config['total_columnas'], *config['encabezados_esperados'].get('columnas_alternativas', [])]:
                problem = 'se esperaban %s columnas; se encontraron %s' % (config['total_columnas'], len(values))
            elif match < 70:
                problem = 'coincidencia de encabezados %.1f%%; se requiere al menos 70%%' % match
            elif words and not keyword_match:
                problem = 'faltan palabras clave del encabezado configurado'
            elif len(mapped) != len([h for h in headers if h]):
                problem = 'hay encabezados duplicados o equivalentes'
            else:
                required = DETAIL_REQUIRED if kind in DETAIL_TYPES else LIST_REQUIRED if kind == 'listado' else set()
                missing = required - mapped
                if missing:
                    problem = 'configure el mapeo de columnas para: %s' % ', '.join(sorted(missing))
                if kind == 'catalogo' and not ({'clave_cucop'} <= mapped and ('codigo_sai' in mapped or 'partida_especifica' in mapped)):
                    problem = 'configure el mapeo de partida específica/clave SAI y clave CUCoP+'
            if problem:
                errors.append('%s: %s.' % (config['name'], problem))
                continue
            candidates.append((config.get('secuencia', 10), not filename_match, config['codigo'], config, ws, headers, match))
        if not candidates:
            raise ImportValidationError('No se detectó un tipo de archivo activo compatible. ' + ' '.join(errors))
        _, _, _, self.config, self.ws, self.headers, self.header_match = min(candidates, key=lambda c: c[:3])
        self.header_row = self.config['fila_encabezado']
        self.subtipo = self.config['tipo_dato']
        self.tipo = 'detalle' if self.subtipo in DETAIL_TYPES else self.subtipo

    def close(self):
        self.wb.close()

    def detect_identifier_from_filename(self):
        match = IDENTIFIER.search(self.filename)
        return match.group(1).upper() if match else None

    def rows(self):
        result = []
        seen = {}
        for idx, row in enumerate(self.ws.iter_rows(min_row=self.header_row + 1, values_only=True), self.header_row + 1):
            if idx - self.header_row > self.MAX_ROWS:
                raise ImportValidationError('Máximo 50,000 renglones por archivo.')
            if all(v is None or v == '' for v in row):
                continue
            values = {k: v for k, v in zip(self.headers, row) if k}
            try:
                if self.tipo == 'anexo':
                    raise ImportValidationError('Anexo reconocido. Su procesamiento no está definido en DEV-02.')
                parsed = self._partida(values) if self.tipo == 'detalle' else self._sai(values) if self.tipo == 'catalogo' else self._procedimiento(values)
                if self.tipo == 'detalle' and values.get('identificador'):
                    identifier_parts(values['identificador'])
                    self.identificadores_contenido.add(code(values['identificador']).upper())
                key = business_key(parsed) if self.tipo == 'detalle' else (parsed['code'], parsed['cucop_code']) if self.tipo == 'catalogo' else parsed['identificador']
                if key in seen:
                    if seen[key] != parsed:
                        raise ImportValidationError('La misma clave de negocio tiene datos contradictorios.')
                    continue
                seen[key] = parsed
                result.append(parsed)
            except ImportValidationError as exc:
                raise ImportValidationError('Fila %s: %s' % (idx, exc)) from exc
        if not result:
            raise ImportValidationError('El archivo no contiene registros. No se marcarán ausencias con un archivo vacío.')
        return result

    def _partida(self, data):
        n = number(data.get('numero'), required=True)
        if n != int(n) or n <= 0:
            raise ImportValidationError('Núm. debe ser un entero positivo.')
        row = {k: code(data.get(k)) for k in PARTIDA_FIELDS if k not in ('numero', 'cantidad', 'cantidad_min', 'cantidad_max', 'cantidad_pendiente')}
        # Deliberately do not strip descriptions: exact business key, including line breaks.
        for k in ('descripcion_cucop', 'descripcion_detallada'):
            row[k] = str(data[k]) if data.get(k) is not None else ''
        if any(not row[k].strip() for k in ('partida_especifica', 'clave_cucop', 'descripcion_detallada', 'unidad_medida')):
            raise ImportValidationError('Complete partida, clave CUCoP+, descripción y unidad de medida.')
        row.update(numero=int(n), cantidad_min=number(data.get('cantidad_min')), cantidad_max=number(data.get('cantidad_max')))
        if 'cantidad_max' in self.headers and row['cantidad_min'] > row['cantidad_max']:
            raise ImportValidationError('Cantidad mínima mayor que la máxima.')
        row['cantidad_pendiente'] = self.subtipo == 'detalle_servicios'
        row['cantidad'] = (0.0 if row['cantidad_pendiente'] else
                           number(data.get('cantidad') if data.get('cantidad') not in (None, '') else data.get('cantidad_max'), required=True))
        return row

    def _procedimiento(self, data):
        identifier = code(data.get('identificador')).upper()
        identifier_parts(identifier)
        row = {k: code(data.get(k)) for k in PORTAL_FIELDS if k not in DATE_FIELDS}
        row['identificador'] = identifier
        row['unidad_codigo'] = row['unidad_codigo'] or identifier.split('-')[3]
        row['tipo_codigo'] = TYPE_CODES.get(norm(data.get('tipo_codigo')))
        if not row['tipo_codigo']:
            raise ImportValidationError('Tipo de contratación no reconocido: %s' % data.get('tipo_codigo'))
        for k in ('nombre_publicado', 'unidad_nombre', 'estatus'):
            if not row[k]:
                raise ImportValidationError('Falta %s.' % k)
        if row['entidad_codigo']:
            row['entidad_codigo'] = row['entidad_codigo'].zfill(2)
            if row['entidad_codigo'] not in {str(n).zfill(2) for n in range(1, 33)}:
                raise ImportValidationError('Código INEGI fuera de 01–32.')
        for optional in ('entidad_codigo', 'codigo_expediente'):
            if optional not in self.headers:
                row.pop(optional)
        for optional in LIST_METADATA_FIELDS:
            if optional not in self.headers:
                row.pop(optional)
        if 'numero_listado' in row:
            value = number(data.get('numero_listado'), required=True)
            if value != int(value) or value <= 0:
                raise ImportValidationError('Núm. debe ser un entero positivo.')
            row['numero_listado'] = int(value)
        # Only present columns may update dates; omitted headers do not erase dates.
        row.update({k: portal_datetime(data.get(k), self.tz_name) for k in DATE_FIELDS if k in self.headers})
        return row

    def _sai(self, data):
        partida = code(data.get('codigo_sai') or data.get('partida_especifica'))
        cucop = code(data.get('clave_cucop'))
        if not partida or not cucop:
            raise ImportValidationError('El catálogo SAI requiere partida específica y clave CUCoP+.')
        return {'code': partida, 'name': code(data.get('nombre_sai')) or partida,
                'cucop_code': cucop, 'cucop_name': code(data.get('descripcion_cucop')) or cucop}
