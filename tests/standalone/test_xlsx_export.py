import importlib.util
import io
import unittest
from datetime import date
from pathlib import Path

import openpyxl
import xlsxwriter

spec = importlib.util.spec_from_file_location('lp_xlsx', Path(__file__).resolve().parents[2] / 'xlsx_export.py')
export = importlib.util.module_from_spec(spec)
spec.loader.exec_module(export)


def render(function, name, rows):
    out = io.BytesIO()
    with xlsxwriter.Workbook(out, {'in_memory': True}) as wb:
        function(wb, name, rows)
    return openpyxl.load_workbook(io.BytesIO(out.getvalue()), data_only=False)


class TestExports(unittest.TestCase):
    def test_costeo_values_and_format(self):
        rows = [['Empresa A', 'PROC', 1, '21601-0028', 'Fibra', 4, 100, 125, 400, 500, 20, 0, 'MXN']]
        wb = render(export.costeo, 'Costeo', rows)
        sheet = wb['Costeo']
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual(sheet['I2'].value, 400)
        self.assertEqual(sheet['J2'].value, 500)
        self.assertEqual(sheet['K2'].value, 20)
        self.assertEqual(sheet.freeze_panes, 'A2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:M2')
        self.assertEqual(sheet['G2'].number_format, '#,##0.00')

    def test_partidas_formula_injection_is_string(self):
        row = ['PROC', 1, '21601', '21601-0028', '=HYPERLINK("https://example.com")', '+SUM(1,2)', 'PIEZA', 0, 0, 100, 'Principal', 'Secundaria', 'Proveedor', 'Participamos', '', True]
        sheet = render(export.partidas, 'Partidas', [row]).active
        self.assertEqual(sheet['E2'].data_type, 's')
        self.assertEqual(sheet['F2'].data_type, 's')
        self.assertEqual(sheet['H2'].value, 0)
        self.assertEqual(sheet['P2'].value, 'Sí')

    def test_checklist_dates_and_empty_values(self):
        row = ['Empresa', 'EXP', 'Proveedor', None, 'Ficha', 'Recibido', date(2026, 9, 21), '', False, 'ficha.pdf']
        sheet = render(export.checklist, 'Checklist', [row]).active
        self.assertEqual(sheet['G2'].number_format, 'dd/mm/yyyy')
        self.assertEqual(sheet['I2'].value, 'No')

    def test_empty_export_has_headers(self):
        for function, name in [(export.costeo, 'Costeo'), (export.partidas, 'Partidas'), (export.checklist, 'Checklist')]:
            with self.subTest(name=name):
                sheet = render(function, name, []).active
                self.assertEqual(sheet.max_row, 1)
                self.assertTrue(sheet['A1'].value)
