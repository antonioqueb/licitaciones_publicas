import copy
import io
import json
import unittest

import openpyxl

from test_parser import parser, profiles, workbook, HEADERS, ROW, IDENTIFIER


class TestDynamicProfiles(unittest.TestCase):
    def reader(self, data=None, name=None, configs=None):
        reader = parser.WorkbookReader(data or workbook(), name or IDENTIFIER + '.xlsx', profiles=configs or profiles())
        self.addCleanup(reader.close)
        return reader

    def test_six_seed_formats_and_no_fallback_without_catalog(self):
        self.assertEqual(len(profiles()), 6)
        with self.assertRaises(parser.ImportValidationError):
            parser.WorkbookReader(workbook(), IDENTIFIER + '.xlsx')

    def test_headers_override_national_character_filename(self):
        services = workbook(HEADERS[:-1], [ROW[:-1]])
        self.assertEqual(self.reader(services).subtipo, 'detalle_servicios')
        self.assertEqual(self.reader(name=IDENTIFIER.replace('-N-', '-T-') + '.xlsx').subtipo, 'detalle_bienes')

    def test_equal_min_max_retains_range(self):
        reader = self.reader(workbook(HEADERS[:-1] + ['Cantidad mínima', 'Cantidad máxima'], [ROW[:-1] + [3, 3]]))
        self.assertEqual(reader.subtipo, 'detalle_rangos')
        row = reader.rows()[0]
        self.assertEqual((row['cantidad_min'], row['cantidad_max'], row['cantidad']), (3, 3, 3))

    def test_range_precedes_requested_quantity(self):
        cfg = profiles()
        next(c for c in cfg if c['tipo_dato'] == 'detalle_rangos')['total_columnas'] = 9
        self.assertEqual(self.reader(workbook(HEADERS + ['Cantidad mínima', 'Cantidad máxima'], [ROW + [2, 4]]), configs=cfg).subtipo, 'detalle_rangos')

    def test_custom_mapping_and_priority_are_dynamic(self):
        cfg = profiles()
        custom = copy.deepcopy(next(c for c in cfg if c['tipo_dato'] == 'detalle_bienes'))
        custom.update(codigo='personalizado', name='Proveedor personalizado', secuencia=1)
        custom['palabras_clave_deteccion'] = '["RUBRO", "CLAVE CUCOP"]'
        mapping = json.loads(custom['mapeo_columnas'])
        mapping['partida_especifica'] = ['Rubro']
        custom['mapeo_columnas'] = json.dumps(mapping)
        expected = json.loads(custom['encabezados_esperados'])
        expected['cols'][1] = 'Rubro'
        custom['encabezados_esperados'] = json.dumps(expected)
        data = workbook([HEADERS[0], 'Rubro', *HEADERS[2:]], [ROW])
        reader = self.reader(data, configs=[custom, *cfg])
        self.assertEqual(reader.config['codigo'], 'personalizado')
        self.assertEqual(reader.rows()[0]['partida_especifica'], '21601')
        custom['activo'] = False
        with self.assertRaises(parser.ImportValidationError):
            self.reader(data, configs=[custom, *cfg])

    def test_priority_resolves_two_valid_profiles(self):
        original = next(c for c in profiles() if c['tipo_dato'] == 'detalle_bienes')
        custom = dict(original, codigo='preferido', secuencia=original['secuencia'] - 1)
        self.assertEqual(self.reader(configs=[original, custom]).config['codigo'], 'preferido')

    def test_seventy_percent_exact_threshold(self):
        cfg = next(c for c in profiles() if c['tipo_dato'] == 'detalle_bienes')
        cfg['encabezados_esperados'] = json.dumps({'cols': [*HEADERS, 'Otro 1', 'Otro 2', 'Otro 3']})
        self.assertEqual(self.reader(configs=[cfg]).header_match, 70)
        cfg['encabezados_esperados'] = json.dumps({'cols': [*HEADERS, 'Otro 1', 'Otro 2', 'Otro 3', 'Otro 4']})
        with self.assertRaisesRegex(parser.ImportValidationError, '70%'):
            self.reader(configs=[cfg])

    def test_sheet_and_header_are_configurable_and_enforced(self):
        cfg = next(c for c in profiles() if c['tipo_dato'] == 'detalle_bienes')
        cfg.update(hojas=2, hoja_nombre='DETALLE MÉDICO', fila_encabezado=3)
        wb = openpyxl.Workbook()
        wb.active.title = 'Portada'
        sheet = wb.create_sheet(' detalle medico ')
        sheet.append(['Título'])
        sheet.append([])
        sheet.append(HEADERS)
        sheet.append(ROW)
        stream = io.BytesIO()
        wb.save(stream)
        self.assertEqual(len(self.reader(stream.getvalue(), configs=[cfg]).rows()), 1)
        cfg['hojas'] = 1
        with self.assertRaisesRegex(parser.ImportValidationError, 'hojas'):
            self.reader(stream.getvalue(), configs=[cfg])
        cfg.update(hojas=2, fila_encabezado=1)
        with self.assertRaises(parser.ImportValidationError):
            self.reader(stream.getvalue(), configs=[cfg])

    def test_json_schema_rejects_expressions_and_ambiguous_mapping(self):
        cfg = profiles()[0]
        for field, invalid in [('palabras_clave_deteccion', "__import__('os')"),
                               ('encabezados_esperados', '{"cols":"incorrecto"}'),
                               ('mapeo_columnas', '{"state":["Estado"]}'),
                               ('mapeo_columnas', '{"identificador":["Uno"],"numero":["UNO"]}')]:
            with self.subTest(field=field, invalid=invalid), self.assertRaises(parser.ImportValidationError):
                parser.profile_config(dict(cfg, **{field: invalid}))

    def test_catalog_and_anexo_detection(self):
        catalog = self.reader(workbook(['Partida específica', 'Clave CUCoP+', 'Descripción SAI', *[f'Otro {i}' for i in range(7)]],
                                      [[21601, '21601-0028', 'Limpieza']]), name='Libro1.xlsx')
        self.assertEqual(catalog.tipo, 'catalogo')
        self.assertEqual(catalog.rows()[0]['code'], '21601')
        wb = openpyxl.Workbook()
        wb.active.title = 'Anexo 1'
        for _ in range(4): wb.active.append([])
        wb.active.append([f'Columna {i}' for i in range(25)])
        wb.active.append(['Dato'])
        stream = io.BytesIO()
        wb.save(stream)
        anexo = self.reader(stream.getvalue(), name='apendice 1.xlsx')
        self.assertEqual(anexo.tipo, 'anexo')
        with self.assertRaisesRegex(parser.ImportValidationError, 'no está definido'):
            anexo.rows()

    def test_filename_accent_and_header_normalization(self):
        self.assertEqual(parser.header_norm('  PARTÍDA   específica  '), 'PARTIDA ESPECIFICA')
        self.assertEqual(parser.filename_norm('InformaciónPública_export_1.xlsx'), 'INFORMACIONPUBLICA_EXPORT_1.XLSX')


if __name__ == '__main__':
    unittest.main()
