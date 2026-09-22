import io

import openpyxl
from odoo.tests import tagged
from odoo.tests.common import HttpCase

from .common import LicitacionCase


@tagged('post_install', '-at_install')
class TestReports(LicitacionCase, HttpCase):
    # HttpCase habilita solicitudes de wkhtmltopdf con el cursor de prueba y
    # libera el bloqueo HTTP mientras se renderiza. No requiere abrir un navegador.
    def test_all_three_pdf_reports(self):
        p, line, exps = self.full_expedientes()
        exp = exps[0]
        self.env['licitacion.pregunta'].create({'expediente_id': exp.id, 'texto': '¿Se acepta ficha técnica equivalente?', 'state': 'lista'})
        for xmlid, docs in [('report_carta_apoyo', exp.carta_apoyo_ids), ('report_expediente', exp), ('report_preguntas_junta', exp)]:
            with self.subTest(report=xmlid):
                report = self.env.ref('licitaciones_publicas.' + xmlid)
                # Odoo devuelve HTML en modo test salvo que se fuerce el PDF real.
                content, fmt = report.with_context(force_report_rendering=True)._render_qweb_pdf(
                    report.report_name, res_ids=docs.ids)
                self.assertEqual(fmt, 'pdf')
                self.assertTrue(content.startswith(b'%PDF-'))
                self.assertGreater(len(content), 5000)

    def test_all_three_xlsx_reports(self):
        p, line, exps = self.full_expedientes()
        exp = exps.filtered(lambda e: e.company_id == line.empresa_principal_id)
        self.env['licitacion.costeo.linea'].create({'expediente_id': exp.id, 'partida_id': line.id, 'costo_unitario': 80, 'precio_unitario': 100})
        self.env['licitacion.documento'].create({'expediente_id': exp.id, 'tipo': 'propio_otro'})
        for xmlid, docs, name in [('action_report_costeo_xlsx', exp, 'Costeo'), ('action_report_partidas_xlsx', p, 'Partidas'), ('action_report_checklist_xlsx', exp, 'Checklist')]:
            with self.subTest(report=xmlid):
                report = self.env.ref('licitaciones_publicas.' + xmlid)
                content, fmt = report.with_context(active_model=docs._name, active_ids=docs.ids)._render_xlsx(report.report_name, docs.ids, data={})
                self.assertEqual(fmt, 'xlsx')
                workbook = openpyxl.load_workbook(io.BytesIO(content))
                self.assertIn(name, workbook.sheetnames)
                self.assertGreater(workbook[name].max_row, 1)
